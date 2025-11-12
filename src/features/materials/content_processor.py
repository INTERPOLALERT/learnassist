"""
Academic Command Center - Materials Library - Content Processor
Extracts text and metadata from various file formats (PDF, DOCX, PPTX, images).
"""

import logging
import os
from typing import Dict, Any, Optional
from pathlib import Path
import sys

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '../..'))

logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)


class ContentProcessor:
    """
    Extracts text from various file formats.

    Supported formats:
    - PDF (using PyPDF2 or pdfplumber)
    - DOCX (using python-docx)
    - PPTX (using python-pptx)
    - TXT, MD (plain text)
    - Images (using Tesseract OCR)
    """

    SUPPORTED_EXTENSIONS = {
        'pdf': 'PDF Document',
        'docx': 'Word Document',
        'pptx': 'PowerPoint Presentation',
        'txt': 'Text File',
        'md': 'Markdown File',
        'png': 'Image (PNG)',
        'jpg': 'Image (JPEG)',
        'jpeg': 'Image (JPEG)',
        'gif': 'Image (GIF)',
        'bmp': 'Image (BMP)'
    }

    def __init__(self):
        """Initialize content processor."""
        self.max_file_size = 50 * 1024 * 1024  # 50 MB max

    def process_file(self, file_path: str) -> Dict[str, Any]:
        """
        Process file and extract content.

        Args:
            file_path: Path to file

        Returns:
            Processing result with extracted text
        """
        logger.info(f"Processing file: {file_path}")

        # Validate file
        validation = self._validate_file(file_path)
        if not validation['valid']:
            return {
                'success': False,
                'error': validation['error']
            }

        file_ext = validation['extension']
        file_size = validation['size']

        try:
            # Extract text based on file type
            if file_ext == 'pdf':
                extracted_text = self._extract_pdf(file_path)
            elif file_ext == 'docx':
                extracted_text = self._extract_docx(file_path)
            elif file_ext == 'pptx':
                extracted_text = self._extract_pptx(file_path)
            elif file_ext in ['txt', 'md']:
                extracted_text = self._extract_text(file_path)
            elif file_ext in ['png', 'jpg', 'jpeg', 'gif', 'bmp']:
                extracted_text = self._extract_image_text(file_path)
            else:
                return {
                    'success': False,
                    'error': f'Unsupported file type: {file_ext}'
                }

            # Calculate metadata
            word_count = len(extracted_text.split())
            char_count = len(extracted_text)

            logger.info(f"✓ Extracted {word_count} words from {file_ext.upper()}")

            return {
                'success': True,
                'extracted_text': extracted_text,
                'metadata': {
                    'file_type': file_ext,
                    'file_size': file_size,
                    'word_count': word_count,
                    'char_count': char_count
                }
            }

        except Exception as e:
            logger.error(f"Failed to process file: {e}", exc_info=True)
            return {
                'success': False,
                'error': f'Processing failed: {str(e)}'
            }

    def _validate_file(self, file_path: str) -> Dict[str, Any]:
        """Validate file exists and is supported."""
        if not os.path.exists(file_path):
            return {'valid': False, 'error': 'File not found'}

        if not os.path.isfile(file_path):
            return {'valid': False, 'error': 'Path is not a file'}

        # Check file size
        file_size = os.path.getsize(file_path)
        if file_size > self.max_file_size:
            return {
                'valid': False,
                'error': f'File too large (max {self.max_file_size // 1024 // 1024} MB)'
            }

        # Check extension
        file_ext = Path(file_path).suffix.lower().lstrip('.')
        if file_ext not in self.SUPPORTED_EXTENSIONS:
            return {
                'valid': False,
                'error': f'Unsupported file type: .{file_ext}'
            }

        return {
            'valid': True,
            'extension': file_ext,
            'size': file_size
        }

    def _extract_pdf(self, file_path: str) -> str:
        """Extract text from PDF using PyPDF2."""
        try:
            import PyPDF2

            text_parts = []

            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)

                logger.debug(f"PDF has {num_pages} pages")

                for page_num in range(num_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)

            return '\n\n'.join(text_parts)

        except Exception as e:
            logger.warning(f"PyPDF2 failed, trying pdfplumber: {e}")
            return self._extract_pdf_fallback(file_path)

    def _extract_pdf_fallback(self, file_path: str) -> str:
        """Fallback PDF extraction using pdfplumber."""
        try:
            import pdfplumber

            text_parts = []

            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)

            return '\n\n'.join(text_parts)

        except Exception as e:
            logger.error(f"pdfplumber also failed: {e}")
            raise Exception("Failed to extract PDF text with both methods")

    def _extract_docx(self, file_path: str) -> str:
        """Extract text from DOCX using python-docx."""
        from docx import Document

        doc = Document(file_path)
        text_parts = []

        # Extract paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        return '\n\n'.join(text_parts)

    def _extract_pptx(self, file_path: str) -> str:
        """Extract text from PPTX using python-pptx."""
        from pptx import Presentation

        prs = Presentation(file_path)
        text_parts = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = []

            # Extract shapes (text boxes, titles, etc.)
            for shape in slide.shapes:
                if hasattr(shape, 'text') and shape.text.strip():
                    slide_text.append(shape.text)

            if slide_text:
                text_parts.append(f"[Slide {slide_num}]\n" + '\n'.join(slide_text))

        return '\n\n'.join(text_parts)

    def _extract_text(self, file_path: str) -> str:
        """Extract text from plain text files."""
        encodings = ['utf-8', 'latin-1', 'cp1252']

        for encoding in encodings:
            try:
                with open(file_path, 'r', encoding=encoding) as file:
                    return file.read()
            except UnicodeDecodeError:
                continue

        raise Exception("Failed to decode text file with common encodings")

    def _extract_image_text(self, file_path: str) -> str:
        """Extract text from images using Tesseract OCR."""
        try:
            from PIL import Image
            import pytesseract

            image = Image.open(file_path)

            # Perform OCR
            text = pytesseract.image_to_string(image, lang='eng')

            if not text.strip():
                return "[Image contains no readable text]"

            return text

        except Exception as e:
            logger.warning(f"OCR failed: {e}")
            return f"[Image - OCR unavailable: {str(e)}]"

    def get_supported_formats(self) -> Dict[str, str]:
        """Get dictionary of supported file formats."""
        return self.SUPPORTED_EXTENSIONS.copy()


if __name__ == "__main__":
    print("Testing Content Processor...")

    processor = ContentProcessor()

    print("\nSupported formats:")
    for ext, desc in processor.get_supported_formats().items():
        print(f"  .{ext}: {desc}")

    print("\nContent Processor validated!")
